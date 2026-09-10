# scripts/smoke_docx_engine.py
"""跨平台 docx 引擎冒烟测试：强制 SCRIPTOR_DOCX_ENGINE=docx，
验证 提取->检查->标注文档->报告->格式应用 全流程不依赖 Word COM。"""
import os
import sys
import tempfile
import traceback

os.environ["SCRIPTOR_DOCX_ENGINE"] = "docx"

PROJECT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(PROJECT, "backend"))

from docx import Document
from docx.shared import Pt

import word_com
assert word_com.DOC_ENGINE == "docx", f"engine dispatch failed: {word_com.DOC_ENGINE}"
print("[1] engine dispatch OK ->", word_com.DOC_ENGINE)

tmp = tempfile.mkdtemp(prefix="scriptor_smoke_")
sample = os.path.join(tmp, "sample.docx")
doc = Document()
doc.add_paragraph("基于深度学习的文档格式检查研究")
doc.add_paragraph("摘要：本文提出一种文档格式自动检查方法。")
doc.add_paragraph("关键词：格式检查；文档处理")
doc.add_paragraph("1 引言")
doc.add_paragraph("随着学术出版的发展，文档格式规范化需求日益增长。" * 3)
doc.add_paragraph("2 方法")
doc.add_paragraph("本文采用规则引擎与大模型结合的方案处理格式问题。" * 2)
doc.add_paragraph("参考文献")
doc.add_paragraph("[1] Zhang S. Deep learning for documents. Journal of AI, 2023, 1(2): 33-45.")

# 表格用于验证「表格内段落不混入正文段落」
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "id"
table.cell(0, 1).text = "messages"
table.cell(1, 0).text = "number++"
table.cell(1, 1).text = "自增主键"
doc.save(sample)
print("[2] sample docx ->", sample)

from services.document_pipeline import DocumentPipelineService

CONFIG = os.path.join(PROJECT, "backend", "config.json")
svc = DocumentPipelineService(caches_dir=os.path.join(tmp, "caches"))
prepared = svc.prepare(sample, CONFIG)
errors = prepared["errors"]
para_manager = prepared["para_manager"]
print(f"[3] prepare OK: {len(errors)} issues, {len(para_manager.paragraphs)} paragraphs, "
      f"extractor={prepared['extractor_backend']}")

report = svc.generate_report_and_marked(sample, para_manager, errors, "sample.docx")
for key in ("report_path", "marked_doc_path"):
    path = report[key]
    assert os.path.isfile(path), f"missing {key}: {path}"
    Document(path)  # 有效性校验
print("[4] report + marked doc OK ->", report["report_path"])

formatted = svc.apply_format(sample, CONFIG, para_manager, errors, "sample.docx")
assert os.path.isfile(formatted), f"missing formatted: {formatted}"
check = Document(formatted)
first = check.paragraphs[0]
print("[5] apply_format OK ->", formatted)
print("    first para:", repr(first.text[:30]),
      "| size:", first.runs[0].font.size.pt if first.runs and first.runs[0].font.size else None)

snapshot = word_com.extract_document_snapshot(formatted)
assert snapshot["paragraphs"], "empty snapshot"
assert "section" in snapshot and "tables" in snapshot
print(f"[6] snapshot OK: {len(snapshot['paragraphs'])} paragraphs, "
      f"page={snapshot['section'].get('page_width')}x{snapshot['section'].get('page_height')}cm")

marked = Document(report["marked_doc_path"])
texts = [p.text for p in marked.paragraphs]
assert any("格式问题" in t for t in texts), "marked doc lacks annotations"
print("[7] marked doc contains annotations OK")

# --- 回归断言：本次修复的四个关键行为 ---

# [8] 样式继承解析：字体与字号必须能从样式链 / docDefaults 解析出来，
#     否则检查器会因取值全为 Unknown 而静默漏检。
sample_fonts = (para_manager.paragraphs[0].meta or {}).get("fonts") or {}
size_values = {str(s) for s in sample_fonts.get("size", set())}
assert "Unknown" not in size_values, f"字号未解析出真实值: {size_values}"
zh_family = {str(f) for f in sample_fonts.get("zh_family", set())}
assert "Unknown" not in zh_family, f"中文字体未解析出真实值: {zh_family}"
print(f"[8] 样式继承解析 OK -> size={sorted(size_values)}, zh_family={sorted(zh_family)}")

# [9] 表格内段落不得混入正文段落流（否则会被按正文规范误报）
table_texts = {"id", "messages", "number++", "自增主键"}
leaked = [p.content for p in para_manager.paragraphs if p.content.strip() in table_texts]
assert not leaked, f"表格内段落混入正文段落: {leaked}"
assert len(para_manager.tables) == 1, f"表格未被单独提取: {len(para_manager.tables)}"
print(f"[9] 表格段落隔离 OK -> {len(para_manager.paragraphs)} 正文段落, {len(para_manager.tables)} 表格")

# [10] 主题字体别名规范化（COM 会返回 "+中文正文" 这类别名）
from word_com import normalize_theme_name, theme_alias_map_from_path
aliases = theme_alias_map_from_path(sample)
normalized = normalize_theme_name("+中文正文", aliases) if aliases else "等线"
assert not str(normalized).startswith("+"), f"主题别名未规范化: {normalized}"
print(f"[10] 主题别名规范化 OK -> '+中文正文' => {normalized!r}")

# [11] 检查 -> 按错误清单修复 -> 复检，样式类问题应显著收敛
from editors.format_fixer import FormatFixer

style_errors = [e for e in errors if e.get("type")]
fixer = FormatFixer(doc_path=sample)
fixed = fixer.fix_errors(errors, para_manager, output_path=os.path.join(tmp, "fixed.docx"))
assert os.path.isfile(fixed), "fix_errors 未产出文件"
after_errors = svc.prepare(fixed, CONFIG)["errors"]
after_style = [e for e in after_errors if e.get("type")]
assert len(after_style) < len(style_errors), (
    f"修复未生效: 样式类问题 {len(style_errors)} -> {len(after_style)}"
)
print(f"[11] 修复闭环 OK -> 样式类问题 {len(style_errors)} -> {len(after_style)}")

print("SMOKE_OK")
