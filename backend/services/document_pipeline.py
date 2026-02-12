from __future__ import annotations

import os
import uuid
from datetime import datetime
from typing import Any, Dict, List, Optional

from docx import Document

from checkers.format_checker import FormatChecker
from editors.document_marker import mark_document_errors
from editors.format_editor import generate_formatted_doc, load_config
from preparation import docx_parser


class DocumentPipelineService:
    def __init__(self, caches_dir: str):
        self.caches_dir = caches_dir
        self.format_checker = FormatChecker()
        os.makedirs(self.caches_dir, exist_ok=True)

    def prepare(
        self,
        doc_path: str,
        config_path: str,
        format_agent: Optional[Any] = None,
    ) -> Dict[str, Any]:
        errors, para_manager = self.format_checker.analyze_format_issues(doc_path, config_path, format_agent)
        doc_content = docx_parser.extract_doc_content(doc_path)
        extractor_backend = self._extract_extractor_backend(para_manager)
        return {
            "errors": errors or [],
            "para_manager": para_manager,
            "doc_content": doc_content,
            "extractor_backend": extractor_backend,
        }

    def generate_report_and_marked(
        self,
        doc_path: str,
        para_manager: Any,
        errors: List[Dict[str, Any]],
        original_filename: str,
    ) -> Dict[str, str]:
        safe_name = self._safe_name(original_filename or os.path.basename(doc_path))
        suffix = uuid.uuid4().hex[:10]
        marked_doc_path = os.path.join(self.caches_dir, f"marked_{safe_name}_{suffix}.docx")
        report_path = os.path.join(self.caches_dir, f"report_{safe_name}_{suffix}.docx")

        mark_document_errors(doc_path, errors, para_manager, marked_doc_path)
        self._build_report_docx(report_path, doc_path, errors)

        return {
            "report_path": report_path,
            "marked_doc_path": marked_doc_path,
        }

    def apply_format(
        self,
        doc_path: str,
        config_path: str,
        para_manager: Any,
        errors: List[Dict[str, Any]],
        original_filename: str,
    ) -> str:
        safe_name = self._safe_name(original_filename or os.path.basename(doc_path))
        suffix = uuid.uuid4().hex[:10]
        output_path = os.path.join(self.caches_dir, f"formatted_{safe_name}_{suffix}.docx")
        config = load_config(config_path)
        return generate_formatted_doc(config, para_manager, output_path, errors or [], doc_path=doc_path)

    @staticmethod
    def _build_report_docx(report_path: str, doc_path: str, errors: List[Dict[str, Any]]) -> None:
        report = Document()
        report.add_heading("Document Format Analysis Report", level=1)
        report.add_paragraph(f"Generated at: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        report.add_paragraph(f"Source document: {os.path.basename(doc_path)}")
        report.add_paragraph(f"Total issues: {len(errors)}")
        report.add_paragraph("")

        if errors:
            for idx, error in enumerate(errors, start=1):
                message = str(error.get("message", "Unknown issue"))
                location = str(error.get("location", ""))
                line = f"{idx}. {message}"
                if location:
                    line += f" (location: {location})"
                report.add_paragraph(line)
        else:
            report.add_paragraph("No format issues were found.")

        report.save(report_path)

    @staticmethod
    def _extract_extractor_backend(para_manager: Any) -> Optional[str]:
        paragraphs = getattr(para_manager, "paragraphs", None)
        if not paragraphs:
            return None
        meta = getattr(paragraphs[0], "meta", {}) or {}
        return meta.get("extractor_backend")

    @staticmethod
    def _safe_name(filename: str) -> str:
        name = os.path.splitext(os.path.basename(filename))[0]
        safe = "".join(ch for ch in name if ch.isalnum() or ch in "-_")
        return safe or "document"
