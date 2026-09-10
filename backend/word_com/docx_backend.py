"""基于 python-docx 的跨平台 Word 处理引擎。

提供与 ``com_utils`` 相同的门面函数与 COM 风格对象模型（``Paragraphs(i)``、
``para.Range.Font``、``para.Range.ParagraphFormat`` 等），使 ``editors/`` 与
``checkers/`` 的代码无需改动即可在 macOS / Linux 上运行。

精度说明：字体（含中文 eastAsia）、字号、加粗、斜体、颜色、对齐、行距、
缩进与段距均通过 python-docx 落盘，不依赖本机安装 Microsoft Word。
由 ``word_com/__init__.py`` 按平台或环境变量 ``SCRIPTOR_DOCX_ENGINE`` 分发。
"""

from __future__ import annotations

import os
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional

from docx import Document as _PyDocxDocument
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Length, Pt, RGBColor

from .com_utils import (
    WD_ALIGN_CENTER,
    WD_ALIGN_JUSTIFY,
    WD_ALIGN_LEFT,
    WD_ALIGN_RIGHT,
    WD_DO_NOT_SAVE_CHANGES,
    WD_FORMAT_DOCX,
    WD_LINE_SPACE_1_5,
    WD_LINE_SPACE_DOUBLE,
    WD_LINE_SPACE_EXACTLY,
    WD_LINE_SPACE_MULTIPLE,
    WD_LINE_SPACE_SINGLE,
    _alignment_to_name,
    _line_spacing_to_string,
    _to_bool_word_flag,
    clean_word_text,
    cm_to_points,
    hex_to_ole_color,
    ole_color_to_hex,
    points_to_cm,
)

from ._docx_styles import (
    ThemeFonts,
    resolve_alignment_code,
    resolve_fonts,
    resolve_indents,
    resolve_line_spacing,
)

_ALIGNMENT_ENUM = {
    int(WD_ALIGN_LEFT): WD_ALIGN_PARAGRAPH.LEFT,
    int(WD_ALIGN_CENTER): WD_ALIGN_PARAGRAPH.CENTER,
    int(WD_ALIGN_RIGHT): WD_ALIGN_PARAGRAPH.RIGHT,
    int(WD_ALIGN_JUSTIFY): WD_ALIGN_PARAGRAPH.JUSTIFY,
}

_RULE_ENUM = {
    int(WD_LINE_SPACE_SINGLE): WD_LINE_SPACING.SINGLE,
    int(WD_LINE_SPACE_1_5): WD_LINE_SPACING.ONE_POINT_FIVE,
    int(WD_LINE_SPACE_DOUBLE): WD_LINE_SPACING.DOUBLE,
    int(WD_LINE_SPACE_EXACTLY): WD_LINE_SPACING.EXACTLY,
    int(WD_LINE_SPACE_MULTIPLE): WD_LINE_SPACING.MULTIPLE,
}


def _rule_int(value: Any) -> int:
    if value is None:
        return int(WD_LINE_SPACE_SINGLE)
    try:
        return int(value)
    except (TypeError, ValueError):
        return int(WD_LINE_SPACE_SINGLE)


def _pt(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, Length):
        return float(value.pt)
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _rfonts(run: Any):
    """获取（必要时创建）``rFonts`` —— 仅用于写入路径。"""
    rPr = run._element.get_or_add_rPr()
    return rPr.get_or_add_rFonts()


def _find_rfonts(run: Any):
    """只读获取 ``rFonts``，不存在时返回 ``None``。

    读取路径必须用它：``get_or_add_rFonts()`` 会向文档注入空节点，
    在只提取格式的场景下会污染原文件 XML。
    """
    try:
        rPr = run._element.rPr
    except Exception:
        return None
    if rPr is None:
        return None
    try:
        return rPr.find(qn("w:rFonts"))
    except Exception:
        return None


def _outline_level_of(para: Any) -> int:
    try:
        style_name = (para.style.name or "").lower() if para.style else ""
    except Exception:
        style_name = ""
    if "heading 1" in style_name or "标题 1" in style_name:
        return 1
    if "heading 2" in style_name or "标题 2" in style_name:
        return 2
    if "heading 3" in style_name or "标题 3" in style_name:
        return 3
    return 10


class _Font:
    """COM ``Range.Font`` 的 python-docx 实现，作用于段落内的全部 run。"""

    def __init__(self, para: Any):
        self._para = para

    def _targets(self) -> List[Any]:
        runs = list(self._para.runs)
        if not runs:
            self._para.add_run("")
            runs = list(self._para.runs)
        return runs

    def _first_run(self) -> Optional[Any]:
        runs = self._para.runs
        return runs[0] if runs else None

    # --- Name ---
    @property
    def Name(self) -> str:
        run = self._first_run()
        return str(run.font.name or "") if run else ""

    @Name.setter
    def Name(self, value: str) -> None:
        for run in self._targets():
            run.font.name = str(value)

    @property
    def NameFarEast(self) -> str:
        run = self._first_run()
        if not run:
            return ""
        fonts = _find_rfonts(run)
        if fonts is None:
            return ""
        return str(fonts.get(qn("w:eastAsia")) or "")

    @NameFarEast.setter
    def NameFarEast(self, value: str) -> None:
        for run in self._targets():
            _rfonts(run).set(qn("w:eastAsia"), str(value))

    @property
    def NameAscii(self) -> str:
        run = self._first_run()
        if not run:
            return ""
        fonts = _find_rfonts(run)
        if fonts is None:
            return ""
        return str(fonts.get(qn("w:ascii")) or "")

    @NameAscii.setter
    def NameAscii(self, value: str) -> None:
        for run in self._targets():
            fonts = _rfonts(run)
            fonts.set(qn("w:ascii"), str(value))
            fonts.set(qn("w:hAnsi"), str(value))

    # --- Size / Bold / Italic / AllCaps ---
    @property
    def Size(self) -> float:
        run = self._first_run()
        return _pt(run.font.size) if run else 0.0

    @Size.setter
    def Size(self, value: Any) -> None:
        size = _pt(value)
        if size <= 0:
            return
        for run in self._targets():
            run.font.size = Pt(size)

    @property
    def Bold(self) -> int:
        run = self._first_run()
        return -1 if (run and run.font.bold) else 0

    @Bold.setter
    def Bold(self, value: Any) -> None:
        flag = value == -1 or bool(value)
        for run in self._targets():
            run.font.bold = flag

    @property
    def Italic(self) -> int:
        run = self._first_run()
        return -1 if (run and run.font.italic) else 0

    @Italic.setter
    def Italic(self, value: Any) -> None:
        flag = value == -1 or bool(value)
        for run in self._targets():
            run.font.italic = flag

    @property
    def AllCaps(self) -> int:
        run = self._first_run()
        return -1 if (run and run.font.all_caps) else 0

    @AllCaps.setter
    def AllCaps(self, value: Any) -> None:
        flag = value == -1 or bool(value)
        for run in self._targets():
            run.font.all_caps = flag

    # --- Color（OLE COLORREF：0x00BBGGRR）---
    @property
    def Color(self) -> int:
        run = self._first_run()
        if not run or run.font.color is None or run.font.color.rgb is None:
            return -1
        rgb = run.font.color.rgb
        return (int(rgb[2]) << 16) | (int(rgb[1]) << 8) | int(rgb[0])

    @Color.setter
    def Color(self, value: Any) -> None:
        try:
            ole = int(value)
        except (TypeError, ValueError):
            return
        if ole < 0:
            return
        color = RGBColor(ole & 0xFF, (ole >> 8) & 0xFF, (ole >> 16) & 0xFF)
        for run in self._targets():
            run.font.color.rgb = color


class _ParaFormat:
    """COM ``Range.ParagraphFormat`` 的 python-docx 实现（磅值语义）。"""

    def __init__(self, para: Any):
        self._pf = para.paragraph_format

    @property
    def LineSpacingRule(self) -> int:
        return _rule_int(self._pf.line_spacing_rule)

    @LineSpacingRule.setter
    def LineSpacingRule(self, value: Any) -> None:
        rule = _RULE_ENUM.get(_rule_int(value))
        if rule is None:
            return
        self._pf.line_spacing_rule = rule

    @property
    def LineSpacing(self) -> float:
        spacing = self._pf.line_spacing
        if spacing is None:
            return 0.0
        if isinstance(spacing, Length):
            return _pt(spacing)
        return float(spacing) * 12.0

    @LineSpacing.setter
    def LineSpacing(self, value: Any) -> None:
        pt = _pt(value)
        if pt <= 0:
            return
        rule = _rule_int(self._pf.line_spacing_rule)
        if rule == int(WD_LINE_SPACE_MULTIPLE):
            self._pf.line_spacing = pt / 12.0
        elif rule == int(WD_LINE_SPACE_EXACTLY):
            self._pf.line_spacing = Pt(pt)
        # SINGLE / 1.5 / DOUBLE 由规则本身决定行距，无需再设值

    @property
    def FirstLineIndent(self) -> float:
        return _pt(self._pf.first_line_indent)

    @FirstLineIndent.setter
    def FirstLineIndent(self, value: Any) -> None:
        self._pf.first_line_indent = Pt(_pt(value))

    @property
    def LeftIndent(self) -> float:
        return _pt(self._pf.left_indent)

    @LeftIndent.setter
    def LeftIndent(self, value: Any) -> None:
        self._pf.left_indent = Pt(_pt(value))

    @property
    def RightIndent(self) -> float:
        return _pt(self._pf.right_indent)

    @RightIndent.setter
    def RightIndent(self, value: Any) -> None:
        self._pf.right_indent = Pt(_pt(value))

    @property
    def SpaceBefore(self) -> float:
        return _pt(self._pf.space_before)

    @SpaceBefore.setter
    def SpaceBefore(self, value: Any) -> None:
        self._pf.space_before = Pt(_pt(value))

    @property
    def SpaceAfter(self) -> float:
        return _pt(self._pf.space_after)

    @SpaceAfter.setter
    def SpaceAfter(self, value: Any) -> None:
        self._pf.space_after = Pt(_pt(value))


class _Style:
    def __init__(self, name: str):
        self.NameLocal = name


class _Range:
    """COM ``Range`` 的简化实现：覆盖一个段落的全部内容。"""

    def __init__(self, para: Any):
        self._para = para

    @property
    def Text(self) -> str:
        text = self._para.text or ""
        return text + "\r"

    @property
    def Font(self) -> _Font:
        return _Font(self._para)

    @property
    def ParagraphFormat(self) -> _ParaFormat:
        return _ParaFormat(self._para)

    @property
    def Style(self) -> _Style:
        try:
            name = self._para.style.name if self._para.style else ""
        except Exception:
            name = ""
        return _Style(str(name or ""))

    def InsertAfter(self, text: str) -> "_Range":
        payload = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
        if payload:
            self._para.add_run(payload)
        return self


class _Paragraph:
    def __init__(self, para: Any):
        self._para = para

    @property
    def Range(self) -> _Range:
        return _Range(self._para)

    @property
    def Alignment(self) -> int:
        value = self._para.alignment
        return int(value) if value is not None else int(WD_ALIGN_LEFT)

    @Alignment.setter
    def Alignment(self, value: Any) -> None:
        enum_value = _ALIGNMENT_ENUM.get(_rule_int(value))
        if enum_value is not None:
            self._para.alignment = enum_value

    @property
    def OutlineLevel(self) -> int:
        return _outline_level_of(self._para)


class _Paragraphs:
    def __init__(self, doc: Any):
        self._doc = doc

    @property
    def Count(self) -> int:
        return len(self._doc.paragraphs)

    def __call__(self, index: int) -> _Paragraph:
        return _Paragraph(self._doc.paragraphs[int(index) - 1])

    def __iter__(self):
        for para in self._doc.paragraphs:
            yield _Paragraph(para)


class _Document:
    def __init__(self, doc: Any, path: Optional[str] = None):
        self._doc = doc
        self._path = os.path.abspath(path) if path else None

    @property
    def Paragraphs(self) -> _Paragraphs:
        return _Paragraphs(self._doc)

    def Save(self) -> None:
        if self._path:
            self._doc.save(self._path)

    def SaveAs2(self, path: str, FileFormat: int = WD_FORMAT_DOCX) -> None:
        self._path = os.path.abspath(str(path))
        Path(self._path).parent.mkdir(parents=True, exist_ok=True)
        self._doc.save(self._path)

    def Close(self, SaveChanges: int = WD_DO_NOT_SAVE_CHANGES) -> None:
        # python-docx 无未保存状态概念，保存由 Save/SaveAs2 显式完成
        return None


class _Documents:
    def Add(self) -> _Document:
        return _Document(_PyDocxDocument(), path=None)

    def Open(self, path: str, ReadOnly: bool = False, **kwargs: Any) -> _Document:
        resolved = str(Path(path).resolve())
        if not os.path.exists(resolved):
            raise FileNotFoundError(f"Document not found: {resolved}")
        return _Document(_PyDocxDocument(resolved), path=resolved)


class _Application:
    def __init__(self):
        self.Visible = False
        self.DisplayAlerts = False
        self.Documents = _Documents()


@contextmanager
def word_session(visible: bool = False):
    yield _Application()


@contextmanager
def open_document(doc_path: str, read_only: bool = True, visible: bool = False):
    resolved = str(Path(doc_path).resolve())
    app = _Application()
    doc = app.Documents.Open(resolved, ReadOnly=bool(read_only))
    try:
        yield app, doc
    finally:
        doc.Close(SaveChanges=WD_DO_NOT_SAVE_CHANGES)


def append_paragraph(doc: _Document, text: str) -> _Range:
    para = doc._doc.add_paragraph(clean_word_text(text))
    return _Range(para)


def build_document(output_path: str, paragraphs: Iterable[str]) -> str:
    target = str(Path(output_path).resolve())
    Path(target).parent.mkdir(parents=True, exist_ok=True)
    doc = _PyDocxDocument()
    for text in paragraphs:
        doc.add_paragraph(clean_word_text(text))
    doc.save(target)
    return target


def _empty_font_dict() -> Dict[str, Any]:
    return {
        "zh_family": "Unknown",
        "en_family": "Unknown",
        "size": 0.0,
        "bold": 0,
        "italic": 0,
        "color": None,
    }


def _font_dict(para: Any, doc: Any = None, theme: Optional[ThemeFonts] = None) -> Dict[str, Any]:
    """还原段落首个 run 实际生效的字体信息。

    走完整的继承链（run → 段落样式 → docDefaults）并解析主题字体引用，
    避免只读 run 直接属性导致的 “字体 Unknown / 字号 0” 漏检。
    """
    if doc is None:
        try:
            doc = para.part.document
        except Exception:
            return _empty_font_dict()
    if theme is None:
        theme = ThemeFonts(doc)
    try:
        return resolve_fonts(para, doc, theme)
    except Exception:
        return _empty_font_dict()


def _extract_document_data(doc_path: str) -> Dict[str, Any]:
    doc = _PyDocxDocument(doc_path)
    theme = ThemeFonts(doc)

    paragraphs: List[Dict[str, Any]] = []
    for idx, para in enumerate(doc.paragraphs, start=1):
        text = clean_word_text(para.text)
        alignment_code = resolve_alignment_code(para, doc)
        spacing_info = resolve_line_spacing(para, doc)
        rule = _rule_int(spacing_info["rule"])
        indents = resolve_indents(para, doc)
        paragraphs.append(
            {
                "index": idx,
                "text": text,
                "alignment_code": alignment_code,
                "alignment": _alignment_to_name(alignment_code),
                "outline_level": _outline_level_of(para),
                "style_name": str(para.style.name) if para.style else "",
                "line_spacing_rule": rule,
                "line_spacing": _line_spacing_to_string(rule, spacing_info["spacing"]),
                "first_line_indent_cm": indents["first_line"],
                "left_indent_cm": indents["left"],
                "right_indent_cm": indents["right"],
                "space_before_cm": indents["space_before"],
                "space_after_cm": indents["space_after"],
                "font": _font_dict(para, doc, theme),
            }
        )

    tables: List[List[List[str]]] = []
    for table in doc.tables:
        rows: List[List[str]] = []
        for row in table.rows:
            rows.append([clean_word_text(cell.text) for cell in row.cells])
        tables.append(rows)

    section_info: Dict[str, Any] = {}
    if doc.sections:
        section = doc.sections[0]
        width_cm = points_to_cm(_pt(section.page_width))
        height_cm = points_to_cm(_pt(section.page_height))
        margin_top = points_to_cm(_pt(section.top_margin))
        margin_bottom = points_to_cm(_pt(section.bottom_margin))
        margin_left = points_to_cm(_pt(section.left_margin))
        margin_right = points_to_cm(_pt(section.right_margin))
        header_distance = points_to_cm(_pt(section.header_distance))
        footer_distance = points_to_cm(_pt(section.footer_distance))
        orientation = "Landscape" if section.orientation == WD_ORIENT.LANDSCAPE else "Portrait"
        section_info = {
            "page_width": width_cm,
            "page_height": height_cm,
            "margin_top": margin_top,
            "margin_bottom": margin_bottom,
            "margin_left": margin_left,
            "margin_right": margin_right,
            "margins": {
                "top": margin_top,
                "bottom": margin_bottom,
                "left": margin_left,
                "right": margin_right,
            },
            "header": {
                "top": header_distance,
                "bottom": footer_distance,
            },
            "orientation": orientation,
            "section_type": 0,
        }

    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "section": section_info,
    }


def extract_document_snapshot(doc_path: str, use_cache: bool = True, use_batch: bool = True) -> Dict[str, Any]:
    """与 ``com_utils.extract_document_snapshot`` 输出结构一致的 python-docx 提取。"""
    resolved = str(Path(doc_path).resolve())

    if use_cache:
        try:
            from .document_cache import get_document_cache
            cache = get_document_cache()
            cached = cache.get(resolved)
            if cached is not None:
                return cached
        except Exception:
            pass

    snapshot = _extract_document_data(resolved)

    if use_cache:
        try:
            from .document_cache import get_document_cache
            get_document_cache().put(resolved, snapshot)
        except Exception:
            pass

    return snapshot
